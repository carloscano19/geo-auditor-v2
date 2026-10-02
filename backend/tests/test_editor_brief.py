import pytest
import io
import docx
from fastapi.testclient import TestClient
from unittest.mock import patch, AsyncMock

from main import app, ai_plan_cache, reset_failed_auth_attempts
from config.settings import Settings, get_settings
from src.utils.docx_brief import (
    generate_editor_brief_docx,
    get_top_non_technical_actions,
    sanitize_domain_for_filename,
    EXCLUDED_TECHNICAL_DIMENSIONS,
    EXCLUDED_SUBMETRIC_KEYWORDS,
)
from src.models.schemas import (
    AuditResponse,
    DimensionScore,
    DetectorResult,
    ScoreBreakdown,
    AIFixesResponse,
    LeadParagraphFix,
    AIPlanResponse,
    PlanQuestion,
    PlanOutlineItem,
    PlanTable,
    PlanDataOpportunity,
    PlanNewParagraph,
    PlanInconsistency,
    PlanSourceToCite,
)


@pytest.fixture(autouse=True)
def reset_state():
    reset_failed_auth_attempts()
    ai_plan_cache.clear()
    get_settings.cache_clear()
    yield
    reset_failed_auth_attempts()
    ai_plan_cache.clear()
    get_settings.cache_clear()


def make_sample_audit_result() -> AuditResponse:
    det_infra = DetectorResult(
        dimension="technical_infrastructure",
        score=40.0,
        weight=0.10,
        contribution=4.0,
        breakdown=[
            ScoreBreakdown(
                name="HTTPS",
                raw_score=0.0,
                weight=0.20,
                weighted_score=0.0,
                explanation="No HTTPS",
                recommendations=["Enable HTTPS immediately"],
            ),
            ScoreBreakdown(
                name="Render Speed",
                raw_score=30.0,
                weight=0.20,
                weighted_score=6.0,
                explanation="Slow TTFB",
                recommendations=["Improve server response time"],
            ),
        ],
    )
    det_meta = DetectorResult(
        dimension="metadata_schema",
        score=30.0,
        weight=0.04,
        contribution=1.2,
        breakdown=[
            ScoreBreakdown(
                name="Schema Presence",
                raw_score=0.0,
                weight=0.40,
                weighted_score=0.0,
                explanation="No schema",
                recommendations=["Add Schema.org JSON-LD"],
            )
        ],
    )
    det_aeo = DetectorResult(
        dimension="aeo_structure",
        score=50.0,
        weight=0.12,
        contribution=6.0,
        breakdown=[
            ScoreBreakdown(
                name="Rule of 60 (Answer First)",
                raw_score=20.0,
                weight=0.50,
                weighted_score=10.0,
                explanation="Direct answer missing in lead",
                recommendations=["Provide a clear 40-60 word answer in the lead paragraph."],
            ),
            ScoreBreakdown(
                name="Interrogative H2s",
                raw_score=40.0,
                weight=0.50,
                weighted_score=20.0,
                explanation="Few question headings",
                recommendations=["Formulate headings as clear user questions."],
            ),
        ],
    )
    det_evid = DetectorResult(
        dimension="evidence_density",
        score=60.0,
        weight=0.12,
        contribution=7.2,
        breakdown=[
            ScoreBreakdown(
                name="Numeric Specificity",
                raw_score=30.0,
                weight=0.50,
                weighted_score=15.0,
                explanation="Lacks specific numbers",
                recommendations=["Include specific percentages and benchmark figures."],
            ),
            ScoreBreakdown(
                name="Attribution Signals",
                raw_score=50.0,
                weight=0.50,
                weighted_score=25.0,
                explanation="Few citations",
                recommendations=["Attribute key claims to named sources."],
            ),
        ],
    )
    return AuditResponse(
        url="https://example.com/blog/seo-guide",
        total_score=65.0,
        dimensions=[
            DimensionScore(name="technical_infrastructure", score=40.0, weight=0.10, contribution=4.0, status="red"),
            DimensionScore(name="aeo_structure", score=50.0, weight=0.12, contribution=6.0, status="yellow"),
            DimensionScore(name="evidence_density", score=60.0, weight=0.12, contribution=7.2, status="yellow"),
        ],
        scoring_version="v2.0",
        language="en",
        content_type="guide_blog",
        analysis_time_ms=1200.0,
        recommendations=["Provide a clear answer in the lead", "Add citations"],
        detector_results=[det_infra, det_meta, det_aeo, det_evid],
    )


def test_top_actions_excludes_technical_metrics():
    audit = make_sample_audit_result()
    top_actions = get_top_non_technical_actions(audit.detector_results, max_actions=8)
    
    # Check that no technical dimension or technical submetric keyword is present
    names = [a[0] for a in top_actions]
    for n in names:
        lower_n = n.lower()
        for kw in EXCLUDED_SUBMETRIC_KEYWORDS:
            assert kw not in lower_n, f"Technical keyword '{kw}' found in top action: {n}"
    
    # Check friendly name mapping
    assert "Answer First in Intro" in names
    assert "Question-based Headings" in names
    assert len(top_actions) <= 8


def test_generate_editor_brief_docx_validity():
    audit = make_sample_audit_result()
    ai_fixes = AIFixesResponse(
        json_ld={"@context": "https://schema.org", "@type": "Article", "headline": "SEO Guide"},
        lead_paragraph=LeadParagraphFix(
            original="This is the old intro paragraph.",
            suggested="This is the optimized opening paragraph that answers the topic directly.",
            rationale="Added inverted pyramid structure.",
        ),
        warnings=["Some warning"],
    )
    ai_plan = AIPlanResponse(
        questions_to_answer=[
            PlanQuestion(
                question="What is SEO?",
                draft_answer="SEO stands for search engine optimization.",
                answer_source="page",
                origin="ai",
            ),
            PlanQuestion(
                question="How much does SEO cost?",
                draft_answer="Costs vary by agency and scope.",
                answer_source="needs_info",
                origin="google_paa",
            ),
        ],
        suggested_h2_structure=[
            PlanOutlineItem(h2="What is SEO?", purpose="Define the concept", status="existing"),
            PlanOutlineItem(h2="Key SEO Strategies", purpose="Actionable list", status="new"),
        ],
        suggested_table=PlanTable(
            title="SEO Comparison",
            headers=["Factor", "On-Page", "Off-Page"],
            rows=[["Focus", "Content", "Backlinks"], ["Control", "High", "Medium"]],
        ),
        data_opportunities=[
            PlanDataOpportunity(suggestion="Add industry CTR benchmarks", source_type="official statistics")
        ],
        paragraphs_to_add=[
            PlanNewParagraph(
                target_issue="Missing direct answer",
                suggested_text="SEO directly influences online visibility by structuring content for both search bots and AI retrieval systems.",
                placement="Under introductory section",
            )
        ],
        inconsistencies=[
            PlanInconsistency(
                issue="Different launch years mentioned",
                values=["2018", "2020"],
                suggestion="Unify all launch year mentions to 2018.",
            )
        ],
        sources_to_cite=[
            PlanSourceToCite(
                url="https://searchengineland.com/seo",
                title="Search Engine Land SEO Guide",
                domain="searchengineland.com",
                found_in="Organic top 10",
                why="Industry reference for crawl standards.",
            )
        ],
        combined_schema={"@context": "https://schema.org", "@graph": []},
        warnings=[],
    )

    docx_bytes = generate_editor_brief_docx(audit, ai_fixes=ai_fixes, ai_plan=ai_plan)
    assert docx_bytes and len(docx_bytes) > 0

    # Verify python-docx can open the generated bytes without error
    doc = docx.Document(io.BytesIO(docx_bytes))
    all_text = " ".join([p.text for p in doc.paragraphs])
    
    assert "Content brief" in all_text
    assert "Citation Score: 65/100" in all_text
    assert "Top actions" in all_text
    assert "Suggested opening paragraph" in all_text
    assert "Inconsistencies found on the page" in all_text
    assert "Questions your page should answer" in all_text
    assert "Suggested H2 structure" in all_text
    assert "Suggested table: SEO Comparison" in all_text
    assert "Data that would enrich the text" in all_text
    assert "Sources to cite or link" in all_text
    assert "Paragraphs to add" in all_text
    assert "For the developer" in all_text

    # Verify footer disclaimer in sections
    footer_text = " ".join([p.text for s in doc.sections for p in s.footer.paragraphs])
    assert "AI-generated suggestions. Review before publishing. They do not affect the Citation Score." in footer_text

    # Verify new badge strings in docx text
    assert "Info already on the page — rewrite it as a Q&A" in all_text
    assert "Missing from the page — needs new content" in all_text

    # Verify table was added
    assert len(doc.tables) >= 1


def test_export_brief_endpoint():
    client = TestClient(app)
    audit = make_sample_audit_result()
    payload = {
        "audit_result": audit.model_dump(mode="json"),
        "ai_fixes": None,
        "ai_plan": None,
    }

    # With no access code required (default settings in test)
    with patch("main.get_settings") as mock_settings:
        mock_settings.return_value = Settings(access_code="")
        response = client.post("/api/export/brief", json=payload)
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert "attachment; filename=\"geo-brief-example.com-" in response.headers["content-disposition"]
        assert len(response.content) > 1000

        # With access code required
        mock_settings.return_value = Settings(access_code="secret123")
        # Without header -> 401
        res_401 = client.post("/api/export/brief", json=payload)
        assert res_401.status_code == 401
        assert res_401.json()["detail"] == "Access code required"

        # With valid header -> 200
        res_200 = client.post("/api/export/brief", json=payload, headers={"X-Access-Code": "secret123"})
        assert res_200.status_code == 200
        assert len(res_200.content) > 1000


@pytest.mark.asyncio
async def test_ai_plan_discards_forbidden_questions():
    """Verify that questions mentioning stock price, share price, revenue, etc. are discarded."""
    client = TestClient(app)
    mock_llm_response = {
        "questions_to_answer": [
            {"question": "What is the stock price of Apple?", "draft_answer": "It is trading high.", "answer_source": "page"},
            {"question": "What is the company revenue?", "draft_answer": "Revenue reached 100 billion.", "answer_source": "page"},
            {"question": "Should I invest in this company?", "draft_answer": "Check financial advisor.", "answer_source": "needs_info"},
            {"question": "¿Cuál es la cotización de las acciones?", "draft_answer": "Cotiza a diez.", "answer_source": "page"},
            {"question": "What are the main product features?", "draft_answer": "Features include fast processing.", "answer_source": "page"},
        ],
        "suggested_h2_structure": [],
        "data_opportunities": [],
        "paragraphs_to_add": [],
        "inconsistencies": [],
    }

    ai_context = {
        "url": "https://example.com/product",
        "title": "Product Overview",
        "main_text": "Features include fast processing. The product is lightweight and durable.",
        "language": "en",
        "content_type": "product",
    }

    with patch("main.get_settings") as mock_settings, \
         patch("main.LLMClient.call_chat_completion", new_callable=AsyncMock) as mock_chat:
        mock_settings.return_value = Settings(
            llm_base_url="https://api.openai.com/v1",
            llm_model="gpt-4o-mini",
            llm_api_key="sk-test",
            access_code="",
        )
        mock_chat.return_value = mock_llm_response

        response = client.post("/api/ai/plan", json={"ai_context": ai_context})
        assert response.status_code == 200
        data = response.json()
        questions = [q["question"] for q in data["questions_to_answer"]]
        
        # Only the legitimate question about features should remain
        assert len(questions) == 1
        assert questions[0] == "What are the main product features?"


@pytest.mark.asyncio
async def test_ai_plan_custom_query_and_cache():
    """Verify custom query passes to SERP client and differentiates cache key."""
    client = TestClient(app)
    ai_context = {
        "url": "https://example.com/guide",
        "title": "Guide to Coffee",
        "main_text": "Making good coffee requires fresh beans and hot water.",
        "language": "en",
        "content_type": "guide_blog",
    }

    mock_llm_response = {
        "questions_to_answer": [
            {"question": "How to make coffee?", "draft_answer": "Use fresh beans.", "answer_source": "page"}
        ],
        "suggested_h2_structure": [],
        "data_opportunities": [],
        "paragraphs_to_add": [],
        "inconsistencies": [],
    }

    mock_serp_response = {
        "market": "US/en",
        "people_also_ask": [],
        "related_searches": [],
        "organic": [],
        "ai_overview_sources": [],
    }

    with patch("main.get_settings") as mock_settings, \
         patch("main.LLMClient.call_chat_completion", new_callable=AsyncMock) as mock_chat, \
         patch("main.SerpClient.fetch_serp_live", new_callable=AsyncMock) as mock_serp:
        mock_settings.return_value = Settings(
            llm_base_url="https://api.openai.com/v1",
            llm_model="gpt-4o-mini",
            llm_api_key="sk-test",
            dataforseo_login="login",
            dataforseo_password="password",
            access_code="",
        )
        mock_chat.return_value = mock_llm_response
        mock_serp.return_value = mock_serp_response

        # Call with custom query
        res1 = client.post("/api/ai/plan", json={
            "ai_context": ai_context,
            "query": "best espresso coffee beans",
        })
        assert res1.status_code == 200
        assert mock_serp.call_count == 1
        # Check that SerpClient was called with the exact custom query
        assert mock_serp.call_args[1]["query"] == "best espresso coffee beans"

        # Calling again with DIFFERENT custom query should NOT hit cache, but call SERP again
        res2 = client.post("/api/ai/plan", json={
            "ai_context": ai_context,
            "query": "cold brew coffee maker",
        })
        assert res2.status_code == 200
        assert mock_serp.call_count == 2
        assert mock_serp.call_args[1]["query"] == "cold brew coffee maker"

        # Calling again with SAME custom query should hit cache (call count stays 2)
        res3 = client.post("/api/ai/plan", json={
            "ai_context": ai_context,
            "query": "cold brew coffee maker",
        })
        assert res3.status_code == 200
        assert mock_serp.call_count == 2

"""
GEO-AUDITOR AI - Word (.docx) Editor Brief Generator

Generates a content brief in Microsoft Word format (.docx) tailored for
non-technical editors, including:
1. Title with URL, date, content type, and Citation Score.
2. Top actions: max 8 submetrics with worst impact (dimension_weight * (100 - score)),
   excluding technical submetrics, with friendly names and recommendations.
3. Suggested opening paragraph (if Quick fixes present) with original and suggested.
4. Plan sections in screen order (inconsistencies, questions, H2 structure,
   suggested table, data opportunities, sources to cite, paragraphs to add),
   each with its explanatory note in italics.
5. "For the developer" annex with Quick fixes JSON-LD and Plan combined schema in monospaced font.
6. Footer: "AI-generated suggestions. Review before publishing. They do not affect the Citation Score."
"""

import io
import json
from datetime import datetime, timezone
from typing import Optional, List, Dict, Any, Tuple
from urllib.parse import urlparse

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from src.models.schemas import (
    AuditResponse,
    AIFixesResponse,
    AIPlanResponse,
    AhrefsOffpageResponse,
    AhrefsLinkingPage,
    DetectorResult,
    ScoreBreakdown,
)

# Dimensions excluded from Top Actions (technical, cannot be touched by content editors)
EXCLUDED_TECHNICAL_DIMENSIONS = {
    "technical_infrastructure",
    "metadata_schema",
}

# Submetric keywords to exclude if they appear in other dimensions
EXCLUDED_SUBMETRIC_KEYWORDS = {
    "https",
    "render",
    "crawl",
    "bot",
    "speed",
    "schema",
}

# Mapping of technical submetric names to friendly editor-oriented names
FRIENDLY_SUBMETRIC_NAMES: Dict[str, str] = {
    "Rule of 60 (Answer First)": "Answer First in Intro",
    "Interrogative H2s": "Question-based Headings",
    "Heading Hierarchy": "Heading Structure",
    "Power Lead (Entity in Lead)": "Main Topic in Opening Sentence",
    "Lexical Richness (MTLD)": "Vocabulary Variety",
    "Autonomous Passages": "Self-contained Sections",
    "Complete Sentences": "Complete Sentences",
    "Content Depth": "Content Depth & Substance",
    "Evidence Density": "Facts & Evidence Density",
    "Numeric Specificity": "Specific Data & Figures",
    "Attribution Signals": "Source Attributions",
    "Author Signals": "Author Information",
    "Freshness Signals": "Publication & Update Dates",
    "Formatting Citability": "Scannable Formatting",
    "Links Verifiability": "Source Links & Citations",
    "Direct Answers": "Direct Answers",
    "Query Relevance": "Search Query Relevance",
    "Keyword Coverage": "Topic Coverage",
    "Entity Identification": "Clear Key Entities",
    "Trust Pages": "About & Editorial Policy",
}


def sanitize_domain_for_filename(url: str) -> str:
    """Extract clean domain for filename: e.g. 'carloscanofernandez.com'."""
    try:
        if not url or url == "text-mode":
            return "text-content"
        if not url.startswith(("http://", "https://")):
            url = "https://" + url
        parsed = urlparse(url)
        domain = parsed.netloc.lower().replace("www.", "")
        cleaned = "".join(c if c.isalnum() or c in ".-" else "_" for c in domain)
        return cleaned.strip(".-_") or "site"
    except Exception:
        return "site"


def get_top_non_technical_actions(
    detector_results: List[DetectorResult], max_actions: int = 8
) -> List[Tuple[str, str, float]]:
    """
    Calculate top non-technical submetrics by impact:
    impact = detector.weight * (100.0 - submetric.raw_score).
    Returns list of (friendly_name, recommendation, impact).
    """
    candidates = []

    for det in detector_results:
        dim = getattr(det, "dimension", "")
        if dim in EXCLUDED_TECHNICAL_DIMENSIONS:
            continue

        raw_w = getattr(det, "weight", None)
        det_weight = 0.10 if raw_w is None else float(raw_w)
        breakdowns = getattr(det, "breakdown", []) or []

        for b in breakdowns:
            name = getattr(b, "name", "")
            raw = getattr(b, "raw_score", None)
            raw_score = 100.0 if raw is None else float(raw)

            # Check if name contains technical keywords
            name_lower = name.lower()
            if any(kw in name_lower for kw in EXCLUDED_SUBMETRIC_KEYWORDS):
                continue

            # Only consider submetrics with raw_score < 70 and with at least one recommendation
            if raw_score >= 70.0:
                continue

            recs = getattr(b, "recommendations", []) or []
            valid_recs = [r.strip() for r in recs if isinstance(r, str) and r.strip()]
            if not valid_recs:
                continue
            rec_text = valid_recs[0]

            impact = det_weight * (100.0 - raw_score)
            friendly_name = FRIENDLY_SUBMETRIC_NAMES.get(name, name)
            candidates.append((friendly_name, rec_text, impact))

    # Merge Heading Structure and Text Walls if both are among candidates
    heading_cand = None
    text_walls_cand = None
    for c in candidates:
        if c[0] == "Heading Structure" and heading_cand is None:
            heading_cand = c
        elif c[0] == "Text Walls" and text_walls_cand is None:
            text_walls_cand = c

    if heading_cand is not None and text_walls_cand is not None:
        merged_impact = max(heading_cand[2], text_walls_cand[2])
        merged_rec = heading_cand[1]
        merged_action = ("Add H2 headings to break up the text", merged_rec, merged_impact)
        candidates = [c for c in candidates if c is not heading_cand and c is not text_walls_cand]
        candidates.append(merged_action)

    # Sort descending by impact
    candidates.sort(key=lambda x: x[2], reverse=True)
    return candidates[:max_actions]


def set_cell_background(cell, color_hex: str):
    """Set shading color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in twentieths of a point (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_callout(doc: Document, text: str, title: Optional[str] = None):
    """Add a shaded callout block for notes or disclaimers."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15

    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.bold = True
        run_title.font.size = Pt(10)
        run_title.font.color.rgb = RGBColor(55, 65, 81)

    run = p.add_run(text)
    run.font.size = Pt(9.5)
    run.font.italic = True
    run.font.color.rgb = RGBColor(75, 85, 99)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def generate_editor_brief_docx(
    audit_result: AuditResponse,
    ai_fixes: Optional[AIFixesResponse] = None,
    ai_plan: Optional[AIPlanResponse] = None,
    ahrefs_offpage: Optional[AhrefsOffpageResponse] = None,
) -> bytes:
    """
    Build the complete Word (.docx) document and return as raw bytes.
    """
    doc = Document()

    # Set standard margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 1. Document Title & Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("Content brief")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(17, 24, 39)

    # Metadata subtitle
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    meta_p.paragraph_format.line_spacing = 1.25

    # Content type formatting
    ct_map = {
        "guide_blog": "Guide / Blog",
        "news": "News / Press",
        "review": "Review",
        "product": "Product",
    }
    content_type_str = ct_map.get(audit_result.content_type, audit_result.content_type or "Guide / Blog")
    
    # Date formatting
    date_str = audit_result.analyzed_at.strftime("%B %d, %Y") if hasattr(audit_result.analyzed_at, "strftime") else str(audit_result.analyzed_at)
    score_display = f"Citation Score: {audit_result.total_score:.0f}/100"

    run_url = meta_p.add_run(f"URL: {audit_result.url}\n")
    run_url.font.size = Pt(10)
    run_url.font.color.rgb = RGBColor(75, 85, 99)

    run_meta = meta_p.add_run(f"Date: {date_str}   |   Type: {content_type_str}   |   ")
    run_meta.font.size = Pt(10)
    run_meta.font.color.rgb = RGBColor(75, 85, 99)

    run_score = meta_p.add_run(score_display)
    run_score.font.size = Pt(10.5)
    run_score.bold = True
    if audit_result.total_score >= 80:
        run_score.font.color.rgb = RGBColor(16, 120, 70)
    elif audit_result.total_score >= 50:
        run_score.font.color.rgb = RGBColor(180, 100, 10)
    else:
        run_score.font.color.rgb = RGBColor(190, 30, 30)

    # 2. Top actions (Max 8 non-technical)
    h2_actions = doc.add_heading(level=1)
    h2_actions.paragraph_format.space_before = Pt(12)
    h2_actions.paragraph_format.space_after = Pt(4)
    run_h2 = h2_actions.add_run("Top actions")
    run_h2.font.name = "Arial"
    run_h2.font.size = Pt(16)
    run_h2.bold = True
    run_h2.font.color.rgb = RGBColor(31, 41, 55)

    sub_actions = doc.add_paragraph()
    sub_actions.paragraph_format.space_after = Pt(8)
    run_sub = sub_actions.add_run(
        "Priority content improvements ranked by their direct impact on AI engine citation."
    )
    run_sub.font.italic = True
    run_sub.font.size = Pt(9.5)
    run_sub.font.color.rgb = RGBColor(107, 114, 128)

    top_actions = get_top_non_technical_actions(audit_result.detector_results, max_actions=8)
    if top_actions:
        for friendly_name, rec, _impact in top_actions:
            bullet = doc.add_paragraph(style="List Bullet")
            bullet.paragraph_format.space_before = Pt(2)
            bullet.paragraph_format.space_after = Pt(4)
            bullet.paragraph_format.line_spacing = 1.15
            
            run_name = bullet.add_run(f"{friendly_name}: ")
            run_name.bold = True
            run_name.font.size = Pt(10)
            run_name.font.color.rgb = RGBColor(17, 24, 39)
            
            run_rec = bullet.add_run(rec)
            run_rec.font.size = Pt(10)
            run_rec.font.color.rgb = RGBColor(55, 65, 81)
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run("No high-priority content actions needed! The page content meets all quality thresholds.")
        r.font.size = Pt(10)
        r.font.italic = True

    # 2.1 Off-page signals (Ahrefs) (if Ahrefs data present)
    if ahrefs_offpage:
        h2_offpage = doc.add_heading(level=1)
        h2_offpage.paragraph_format.space_before = Pt(16)
        h2_offpage.paragraph_format.space_after = Pt(4)
        run_h2_off = h2_offpage.add_run("Off-page signals (Ahrefs)")
        run_h2_off.font.name = "Arial"
        run_h2_off.font.size = Pt(16)
        run_h2_off.bold = True
        run_h2_off.font.color.rgb = RGBColor(31, 41, 55)

        sub_off = doc.add_paragraph()
        sub_off.paragraph_format.space_after = Pt(8)
        run_sub_off = sub_off.add_run("Data from Ahrefs. Not included in the Citation Score.")
        run_sub_off.font.italic = True
        run_sub_off.font.size = Pt(9.5)
        run_sub_off.font.color.rgb = RGBColor(107, 114, 128)

        # Off-page Metrics Table
        table_off = doc.add_table(rows=1, cols=2)
        table_off.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_off.autofit = True

        hdr_cells = table_off.rows[0].cells
        hdr_cells[0].text = "Signal"
        hdr_cells[1].text = "Value"
        for c in hdr_cells:
            set_cell_background(c, "F9FAFB")
            set_cell_margins(c, 100, 100, 150, 150)
            for p in c.paragraphs:
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(55, 65, 81)

        def _fmt(v: Any) -> str:
            if v is None:
                return "n/a"
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            if isinstance(v, (int, float)):
                return f"{v:,.0f}" if isinstance(v, int) else f"{v:,.1f}"
            return str(v)

        metrics_list = [
            ("Domain Rating", _fmt(ahrefs_offpage.domain_rating)),
            ("URL Rating", _fmt(ahrefs_offpage.url_rating)),
            ("Referring domains", _fmt(ahrefs_offpage.referring_domains)),
            ("Backlinks", _fmt(ahrefs_offpage.backlinks)),
            ("Organic keywords", _fmt(ahrefs_offpage.organic_keywords)),
            ("Top 3 keywords", _fmt(ahrefs_offpage.top3_keywords)),
            ("Organic traffic/month", _fmt(ahrefs_offpage.organic_traffic)),
        ]

        for label, val_str in metrics_list:
            row_cells = table_off.add_row().cells
            row_cells[0].text = label
            row_cells[1].text = val_str
            for c in row_cells:
                set_cell_margins(c, 80, 80, 150, 150)
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(55, 65, 81)
            for p in row_cells[1].paragraphs:
                for r in p.runs:
                    r.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        if ahrefs_offpage.recommendations:
            p_rec_title = doc.add_paragraph()
            p_rec_title.paragraph_format.space_before = Pt(6)
            p_rec_title.paragraph_format.space_after = Pt(2)
            r_rec_title = p_rec_title.add_run("Off-page recommendations:")
            r_rec_title.bold = True
            r_rec_title.font.size = Pt(10)
            r_rec_title.font.color.rgb = RGBColor(31, 41, 55)

            for rec in ahrefs_offpage.recommendations:
                b_rec = doc.add_paragraph(style="List Bullet")
                b_rec.paragraph_format.space_before = Pt(2)
                b_rec.paragraph_format.space_after = Pt(4)
                r_b = b_rec.add_run(rec)
                r_b.font.size = Pt(9.5)
                r_b.font.color.rgb = RGBColor(55, 65, 81)

        # Backlinks Table ("Who links to this page")
        if ahrefs_offpage.linking_pages:
            p_bl_title = doc.add_paragraph()
            p_bl_title.paragraph_format.space_before = Pt(10)
            p_bl_title.paragraph_format.space_after = Pt(4)
            r_bl_title = p_bl_title.add_run("Who links to this page:")
            r_bl_title.bold = True
            r_bl_title.font.size = Pt(10)
            r_bl_title.font.color.rgb = RGBColor(31, 41, 55)

            table_bl = doc.add_table(rows=1, cols=6)
            table_bl.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_bl.autofit = True

            bl_hdr = table_bl.rows[0].cells
            bl_hdr_titles = ["Domain", "DR", "Linking page", "Anchor", "Type", "First seen"]
            for idx, title in enumerate(bl_hdr_titles):
                bl_hdr[idx].text = title
                set_cell_background(bl_hdr[idx], "F9FAFB")
                set_cell_margins(bl_hdr[idx], 80, 80, 100, 100)
                for p in bl_hdr[idx].paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = RGBColor(55, 65, 81)

            for lp in ahrefs_offpage.linking_pages[:20]:
                r_cells = table_bl.add_row().cells
                is_sp = getattr(lp, "spam", False) if not isinstance(lp, dict) else lp.get("spam", False)
                dom = (getattr(lp, "domain", "") if not isinstance(lp, dict) else lp.get("domain", "")) or ""
                dom_display = f"{dom} [Spam]" if is_sp else dom
                dr_val = getattr(lp, "domain_rating", None) if not isinstance(lp, dict) else lp.get("domain_rating")
                url_f = (getattr(lp, "url_from", "") if not isinstance(lp, dict) else lp.get("url_from", "")) or ""
                anch = (getattr(lp, "anchor", "") if not isinstance(lp, dict) else lp.get("anchor", "")) or "—"
                is_df = getattr(lp, "dofollow", False) if not isinstance(lp, dict) else lp.get("dofollow", False)
                type_str = "Dofollow" if is_df else "Nofollow"
                if is_sp:
                    type_str += " (Spam)"
                fseen = (getattr(lp, "first_seen", "") if not isinstance(lp, dict) else lp.get("first_seen", "")) or "—"

                r_cells[0].text = dom_display
                r_cells[1].text = _fmt(dr_val)
                r_cells[2].text = url_f
                r_cells[3].text = anch
                r_cells[4].text = type_str
                r_cells[5].text = fseen

                for c in r_cells:
                    set_cell_margins(c, 60, 60, 100, 100)
                    for p in c.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(8)
                            r.font.color.rgb = RGBColor(55, 65, 81)

    # 3. Suggested opening paragraph (if Quick fixes present)
    if ai_fixes and ai_fixes.lead_paragraph:
        lp = ai_fixes.lead_paragraph
        h2_lead = doc.add_heading(level=1)
        h2_lead.paragraph_format.space_before = Pt(16)
        h2_lead.paragraph_format.space_after = Pt(4)
        run_h2_lead = h2_lead.add_run("Suggested opening paragraph")
        run_h2_lead.font.name = "Arial"
        run_h2_lead.font.size = Pt(16)
        run_h2_lead.bold = True
        run_h2_lead.font.color.rgb = RGBColor(31, 41, 55)

        sub_lead = doc.add_paragraph()
        sub_lead.paragraph_format.space_after = Pt(8)
        run_sub_lead = sub_lead.add_run(
            "Rewrite the opening paragraph to answer the primary topic immediately (inverted pyramid structure)."
        )
        run_sub_lead.font.italic = True
        run_sub_lead.font.size = Pt(9.5)
        run_sub_lead.font.color.rgb = RGBColor(107, 114, 128)

        if lp.original:
            p_orig = doc.add_paragraph()
            p_orig.paragraph_format.space_before = Pt(4)
            p_orig.paragraph_format.space_after = Pt(2)
            run_lbl_orig = p_orig.add_run("Original opening:\n")
            run_lbl_orig.bold = True
            run_lbl_orig.font.size = Pt(10)
            run_lbl_orig.font.color.rgb = RGBColor(107, 114, 128)
            run_txt_orig = p_orig.add_run(f"\"{lp.original}\"")
            run_txt_orig.font.size = Pt(10)
            run_txt_orig.font.color.rgb = RGBColor(107, 114, 128)

        p_sug = doc.add_paragraph()
        p_sug.paragraph_format.space_before = Pt(4)
        p_sug.paragraph_format.space_after = Pt(4)
        run_lbl_sug = p_sug.add_run("Suggested revision:\n")
        run_lbl_sug.bold = True
        run_lbl_sug.font.size = Pt(10.5)
        run_lbl_sug.font.color.rgb = RGBColor(16, 120, 70)
        run_txt_sug = p_sug.add_run(f"\"{lp.suggested}\"")
        run_txt_sug.font.size = Pt(10.5)
        run_txt_sug.bold = True
        run_txt_sug.font.color.rgb = RGBColor(17, 24, 39)

        if lp.rationale:
            p_rat = doc.add_paragraph()
            p_rat.paragraph_format.space_after = Pt(8)
            run_lbl_rat = p_rat.add_run("Rationale: ")
            run_lbl_rat.bold = True
            run_lbl_rat.font.size = Pt(9.5)
            run_lbl_rat.font.color.rgb = RGBColor(75, 85, 99)
            run_txt_rat = p_rat.add_run(lp.rationale)
            run_txt_rat.font.size = Pt(9.5)
            run_txt_rat.font.italic = True
            run_txt_rat.font.color.rgb = RGBColor(75, 85, 99)

    # 4. Plan Sections (if AI Plan present)
    if ai_plan:
        # 4.1 Inconsistencies found on the page
        if ai_plan.inconsistencies and len(ai_plan.inconsistencies) > 0:
            h2_inc = doc.add_heading(level=1)
            h2_inc.paragraph_format.space_before = Pt(16)
            h2_inc.paragraph_format.space_after = Pt(4)
            run_h2_inc = h2_inc.add_run("Inconsistencies found on the page")
            run_h2_inc.font.name = "Arial"
            run_h2_inc.font.size = Pt(16)
            run_h2_inc.bold = True
            run_h2_inc.font.color.rgb = RGBColor(31, 41, 55)

            sub_inc = doc.add_paragraph()
            sub_inc.paragraph_format.space_after = Pt(8)
            run_sub_inc = sub_inc.add_run(
                "The page contradicts itself here. Pick one version and use it everywhere, so AI engines extract a single answer."
            )
            run_sub_inc.font.italic = True
            run_sub_inc.font.size = Pt(9.5)
            run_sub_inc.font.color.rgb = RGBColor(107, 114, 128)

            for inc in ai_plan.inconsistencies:
                p_inc = doc.add_paragraph(style="List Bullet")
                p_inc.paragraph_format.space_before = Pt(2)
                p_inc.paragraph_format.space_after = Pt(4)
                
                r_issue = p_inc.add_run(f"{inc.issue}\n")
                r_issue.bold = True
                r_issue.font.size = Pt(10)
                
                vals_str = " vs ".join([f"\"{v}\"" for v in inc.values])
                r_vals = p_inc.add_run(f"Conflicting values: {vals_str}\n")
                r_vals.font.size = Pt(9.5)
                r_vals.font.color.rgb = RGBColor(180, 50, 50)
                
                r_sug = p_inc.add_run(f"Recommendation: {inc.suggestion}")
                r_sug.font.size = Pt(9.5)
                r_sug.font.color.rgb = RGBColor(55, 65, 81)

        # 4.2 Questions your page should answer
        if ai_plan.questions_to_answer and len(ai_plan.questions_to_answer) > 0:
            h2_q = doc.add_heading(level=1)
            h2_q.paragraph_format.space_before = Pt(16)
            h2_q.paragraph_format.space_after = Pt(4)
            run_h2_q = h2_q.add_run("Questions your page should answer")
            run_h2_q.font.name = "Arial"
            run_h2_q.font.size = Pt(16)
            run_h2_q.bold = True
            run_h2_q.font.color.rgb = RGBColor(31, 41, 55)

            sub_q = doc.add_paragraph()
            sub_q.paragraph_format.space_after = Pt(8)
            run_sub_q = sub_q.add_run(
                "Add these as an FAQ block. Short, direct answers are what AI engines quote."
            )
            run_sub_q.font.italic = True
            run_sub_q.font.size = Pt(9.5)
            run_sub_q.font.color.rgb = RGBColor(107, 114, 128)

            for q in ai_plan.questions_to_answer:
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(4)
                p_q.paragraph_format.space_after = Pt(2)

                # Badge text
                badge = (
                    "Info already on the page — rewrite it as a Q&A"
                    if q.answer_source == "page"
                    else "Missing from the page — needs new content"
                )
                paa_extra = " · Real Google question" if q.origin == "google_paa" else ""

                r_qtitle = p_q.add_run(f"Q: {q.question}\n")
                r_qtitle.bold = True
                r_qtitle.font.size = Pt(10.5)
                r_qtitle.font.color.rgb = RGBColor(17, 24, 39)

                r_qbadge = p_q.add_run(f"[{badge}{paa_extra}]\n")
                r_qbadge.font.size = Pt(9)
                r_qbadge.bold = True
                if q.answer_source == "page":
                    r_qbadge.font.color.rgb = RGBColor(16, 120, 70)
                else:
                    r_qbadge.font.color.rgb = RGBColor(180, 100, 10)

                r_qans = p_q.add_run(f"A: {q.draft_answer}")
                r_qans.font.size = Pt(10)
                r_qans.font.color.rgb = RGBColor(55, 65, 81)

        # 4.3 Suggested H2 structure
        if ai_plan.suggested_h2_structure and len(ai_plan.suggested_h2_structure) > 0:
            h2_h2 = doc.add_heading(level=1)
            h2_h2.paragraph_format.space_before = Pt(16)
            h2_h2.paragraph_format.space_after = Pt(4)
            run_h2_h2 = h2_h2.add_run("Suggested H2 structure")
            run_h2_h2.font.name = "Arial"
            run_h2_h2.font.size = Pt(16)
            run_h2_h2.bold = True
            run_h2_h2.font.color.rgb = RGBColor(31, 41, 55)

            sub_h2 = doc.add_paragraph()
            sub_h2.paragraph_format.space_after = Pt(8)
            run_sub_h2 = sub_h2.add_run(
                "Use these headings to organise the page. 'New' means the section has to be written."
            )
            run_sub_h2.font.italic = True
            run_sub_h2.font.size = Pt(9.5)
            run_sub_h2.font.color.rgb = RGBColor(107, 114, 128)

            for item in ai_plan.suggested_h2_structure:
                p_item = doc.add_paragraph(style="List Bullet")
                p_item.paragraph_format.space_before = Pt(2)
                p_item.paragraph_format.space_after = Pt(3)

                status_label = "Exists — add the heading" if item.status == "existing" else "New — write this section"
                r_title = p_item.add_run(f"{item.h2} ")
                r_title.bold = True
                r_title.font.size = Pt(10)
                r_title.font.color.rgb = RGBColor(17, 24, 39)

                r_st = p_item.add_run(f"[{status_label}]\n")
                r_st.font.size = Pt(9)
                r_st.bold = True
                r_st.font.color.rgb = RGBColor(100, 100, 180) if item.status == "new" else RGBColor(107, 114, 128)

                r_purp = p_item.add_run(f"Purpose: {item.purpose}")
                r_purp.font.size = Pt(9.5)
                r_purp.font.color.rgb = RGBColor(75, 85, 99)

        # 4.4 Suggested table
        if ai_plan.suggested_table:
            st = ai_plan.suggested_table
            h2_tbl = doc.add_heading(level=1)
            h2_tbl.paragraph_format.space_before = Pt(16)
            h2_tbl.paragraph_format.space_after = Pt(4)
            run_h2_tbl = h2_tbl.add_run(f"Suggested table: {st.title}")
            run_h2_tbl.font.name = "Arial"
            run_h2_tbl.font.size = Pt(16)
            run_h2_tbl.bold = True
            run_h2_tbl.font.color.rgb = RGBColor(31, 41, 55)

            sub_tbl = doc.add_paragraph()
            sub_tbl.paragraph_format.space_after = Pt(8)
            run_sub_tbl = sub_tbl.add_run(
                "Add this table to the page. All values come from the page itself."
            )
            run_sub_tbl.font.italic = True
            run_sub_tbl.font.size = Pt(9.5)
            run_sub_tbl.font.color.rgb = RGBColor(107, 114, 128)

            if st.headers and st.rows and len(st.rows) > 0:
                tbl = doc.add_table(rows=len(st.rows) + 1, cols=len(st.headers))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = True

                # Header row
                hdr_cells = tbl.rows[0].cells
                for i, header_text in enumerate(st.headers):
                    cell = hdr_cells[i]
                    set_cell_background(cell, "E5E7EB")
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    r = p.add_run(str(header_text))
                    r.bold = True
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(17, 24, 39)

                # Data rows
                for r_idx, row_data in enumerate(st.rows):
                    row_cells = tbl.rows[r_idx + 1].cells
                    for c_idx, cell_value in enumerate(row_data):
                        if c_idx < len(row_cells):
                            cell = row_cells[c_idx]
                            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                            # Alternate row shading
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "F9FAFB")
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(0)
                            p.paragraph_format.space_after = Pt(0)
                            r = p.add_run(str(cell_value))
                            r.font.size = Pt(9.5)
                            r.font.color.rgb = RGBColor(55, 65, 81)

                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            elif st.table_idea:
                p_idea = doc.add_paragraph()
                p_idea.paragraph_format.space_after = Pt(6)
                r_lbl = p_idea.add_run("Table Idea: ")
                r_lbl.bold = True
                r_lbl.font.size = Pt(10)
                r_txt = p_idea.add_run(st.table_idea)
                r_txt.font.size = Pt(10)
                r_txt.font.italic = True

        # 4.5 Data that would enrich the text
        if ai_plan.data_opportunities and len(ai_plan.data_opportunities) > 0:
            h2_data = doc.add_heading(level=1)
            h2_data.paragraph_format.space_before = Pt(16)
            h2_data.paragraph_format.space_after = Pt(4)
            run_h2_data = h2_data.add_run("Data that would enrich the text")
            run_h2_data.font.name = "Arial"
            run_h2_data.font.size = Pt(16)
            run_h2_data.bold = True
            run_h2_data.font.color.rgb = RGBColor(31, 41, 55)

            sub_data = doc.add_paragraph()
            sub_data.paragraph_format.space_after = Pt(8)
            run_sub_data = sub_data.add_run(
                "Information worth adding. No figures are suggested: find them in the type of source shown."
            )
            run_sub_data.font.italic = True
            run_sub_data.font.size = Pt(9.5)
            run_sub_data.font.color.rgb = RGBColor(107, 114, 128)

            for d in ai_plan.data_opportunities:
                p_d = doc.add_paragraph(style="List Bullet")
                p_d.paragraph_format.space_before = Pt(2)
                p_d.paragraph_format.space_after = Pt(3)

                r_sug = p_d.add_run(f"{d.suggestion} ")
                r_sug.bold = True
                r_sug.font.size = Pt(10)
                r_sug.font.color.rgb = RGBColor(17, 24, 39)

                r_src = p_d.add_run(f"(Recommended Source: {d.source_type})")
                r_src.font.size = Pt(9.5)
                r_src.font.italic = True
                r_src.font.color.rgb = RGBColor(75, 85, 99)

        # 4.6 Sources to cite or link
        if ai_plan.sources_to_cite and len(ai_plan.sources_to_cite) > 0:
            h2_src = doc.add_heading(level=1)
            h2_src.paragraph_format.space_before = Pt(16)
            h2_src.paragraph_format.space_after = Pt(4)
            run_h2_src = h2_src.add_run("Sources to cite or link")
            run_h2_src.font.name = "Arial"
            run_h2_src.font.size = Pt(16)
            run_h2_src.bold = True
            run_h2_src.font.color.rgb = RGBColor(31, 41, 55)

            sub_src = doc.add_paragraph()
            sub_src.paragraph_format.space_after = Pt(8)
            run_sub_src = sub_src.add_run(
                "Real pages that Google ranks or cites for this topic. Link or cite the relevant ones."
            )
            run_sub_src.font.italic = True
            run_sub_src.font.size = Pt(9.5)
            run_sub_src.font.color.rgb = RGBColor(107, 114, 128)

            for s in ai_plan.sources_to_cite:
                p_s = doc.add_paragraph(style="List Bullet")
                p_s.paragraph_format.space_before = Pt(2)
                p_s.paragraph_format.space_after = Pt(4)

                r_title = p_s.add_run(f"{s.title} ({s.domain})\n")
                r_title.bold = True
                r_title.font.size = Pt(10)
                r_title.font.color.rgb = RGBColor(17, 24, 39)

                r_url = p_s.add_run(f"URL: {s.url}   [{s.found_in}]\n")
                r_url.font.size = Pt(9)
                r_url.font.color.rgb = RGBColor(37, 99, 235)

                if s.why:
                    r_why = p_s.add_run(f"Why cite: {s.why}")
                    r_why.font.size = Pt(9.5)
                    r_why.font.italic = True
                    r_why.font.color.rgb = RGBColor(75, 85, 99)

        # 4.7 Paragraphs to add
        if ai_plan.paragraphs_to_add and len(ai_plan.paragraphs_to_add) > 0:
            h2_p = doc.add_heading(level=1)
            h2_p.paragraph_format.space_before = Pt(16)
            h2_p.paragraph_format.space_after = Pt(4)
            run_h2_p = h2_p.add_run("Paragraphs to add")
            run_h2_p.font.name = "Arial"
            run_h2_p.font.size = Pt(16)
            run_h2_p.bold = True
            run_h2_p.font.color.rgb = RGBColor(31, 41, 55)

            sub_p = doc.add_paragraph()
            sub_p.paragraph_format.space_after = Pt(8)
            run_sub_p = sub_p.add_run(
                "Ready-to-paste paragraphs, written only with facts from the page."
            )
            run_sub_p.font.italic = True
            run_sub_p.font.size = Pt(9.5)
            run_sub_p.font.color.rgb = RGBColor(107, 114, 128)

            for i, p_item in enumerate(ai_plan.paragraphs_to_add, 1):
                p_para = doc.add_paragraph()
                p_para.paragraph_format.space_before = Pt(6)
                p_para.paragraph_format.space_after = Pt(2)

                r_head = p_para.add_run(f"Paragraph {i}: Addressing \"{p_item.target_issue}\"\n")
                r_head.bold = True
                r_head.font.size = Pt(10.5)
                r_head.font.color.rgb = RGBColor(17, 24, 39)

                r_place = p_para.add_run(f"Placement: {p_item.placement}\n")
                r_place.font.size = Pt(9.5)
                r_place.font.italic = True
                r_place.font.color.rgb = RGBColor(107, 114, 128)

                # Shaded text block for the paragraph
                add_callout(doc, p_item.suggested_text)

    # 5. Annex: "For the developer" (if Quick fixes JSON-LD or Plan combined_schema present)
    has_dev_schema = (ai_fixes and ai_fixes.json_ld) or (ai_plan and ai_plan.combined_schema)
    if has_dev_schema:
        h2_dev = doc.add_heading(level=1)
        h2_dev.paragraph_format.space_before = Pt(20)
        h2_dev.paragraph_format.space_after = Pt(4)
        run_h2_dev = h2_dev.add_run("For the developer")
        run_h2_dev.font.name = "Arial"
        run_h2_dev.font.size = Pt(16)
        run_h2_dev.bold = True
        run_h2_dev.font.color.rgb = RGBColor(31, 41, 55)

        sub_dev = doc.add_paragraph()
        sub_dev.paragraph_format.space_after = Pt(8)
        run_sub_dev = sub_dev.add_run(
            "Structured Data (Schema.org JSON-LD) to add inside <script type=\"application/ld+json\"> in the page <head>."
        )
        run_sub_dev.font.italic = True
        run_sub_dev.font.size = Pt(9.5)
        run_sub_dev.font.color.rgb = RGBColor(107, 114, 128)

        if ai_fixes and ai_fixes.json_ld:
            p_lbl_ld = doc.add_paragraph()
            p_lbl_ld.paragraph_format.space_before = Pt(4)
            p_lbl_ld.paragraph_format.space_after = Pt(2)
            r = p_lbl_ld.add_run("Article Schema.org JSON-LD (Quick fixes):")
            r.bold = True
            r.font.size = Pt(10)
            
            p_code = doc.add_paragraph()
            p_code.paragraph_format.space_after = Pt(8)
            r_code = p_code.add_run(json.dumps(ai_fixes.json_ld, indent=2))
            r_code.font.name = "Courier New"
            r_code.font.size = Pt(8.5)
            r_code.font.color.rgb = RGBColor(30, 41, 59)

        if ai_plan and ai_plan.combined_schema:
            p_lbl_cs = doc.add_paragraph()
            p_lbl_cs.paragraph_format.space_before = Pt(4)
            p_lbl_cs.paragraph_format.space_after = Pt(2)
            r = p_lbl_cs.add_run("Combined Schema.org JSON-LD (Full Plan):")
            r.bold = True
            r.font.size = Pt(10)

            p_code = doc.add_paragraph()
            p_code.paragraph_format.space_after = Pt(8)
            r_code = p_code.add_run(json.dumps(ai_plan.combined_schema, indent=2))
            r_code.font.name = "Courier New"
            r_code.font.size = Pt(8.5)
            r_code.font.color.rgb = RGBColor(30, 41, 59)

    # 6. Footer disclaimer
    for sec in doc.sections:
        footer = sec.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_foot.paragraph_format.space_before = Pt(6)
        r_foot = p_foot.add_run(
            "AI-generated suggestions. Review before publishing. They do not affect the Citation Score."
        )
        r_foot.font.name = "Arial"
        r_foot.font.size = Pt(8.5)
        r_foot.font.italic = True
        r_foot.font.color.rgb = RGBColor(156, 163, 175)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
